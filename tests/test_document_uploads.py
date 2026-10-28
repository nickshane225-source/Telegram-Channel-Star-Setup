import io
from pathlib import Path
import pytest
from PIL import Image
from fastapi import UploadFile
import docx
import pypdfium2

from app.services.media import MediaService


@pytest.mark.asyncio
async def test_validate_file_extensions(tmp_path):
    """Verify allowed and disallowed extensions."""
    svc = MediaService(media_dir=tmp_path / "media")

    for allowed in ["photo.jpg", "pic.png", "img.webp", "doc.pdf", "file.docx", "old.doc", "notes.txt"]:
        ok, err = svc.validate_file(allowed, 1024)
        assert ok is True, f"Expected {allowed} to be valid, got {err}"

    ok, err = svc.validate_file("malicious.exe", 1024)
    assert ok is False
    assert "Unsupported file extension" in err


@pytest.mark.asyncio
async def test_process_txt_upload(tmp_path):
    """Verify text file (.txt) is converted into formatted publication slide cards."""
    svc = MediaService(media_dir=tmp_path / "media")

    text_content = (
        "Project Architecture & Telegram Stars Integration\n\n"
        "This is an automated text document card generated for Telegram Stars Paid Media.\n"
        "It supports UTF-8 characters and multi-line formatting."
    ).encode("utf-8")

    upload = UploadFile(filename="briefing.txt", file=io.BytesIO(text_content))
    photos = await svc.save_and_process_file_upload(upload)

    assert len(photos) >= 1
    p = photos[0]
    assert Path(p.path).exists()
    assert p.source_type == "txt"
    assert p.width == 1200
    assert p.height == 1600
    assert p.size_bytes > 0


@pytest.mark.asyncio
async def test_process_docx_upload(tmp_path):
    """Verify Word (.docx) document is parsed and rendered into styled cards."""
    svc = MediaService(media_dir=tmp_path / "media")

    # Build docx in memory
    doc = docx.Document()
    doc.add_heading("Release Notes v2.0", level=0)
    doc.add_paragraph("This document contains exclusive paid content for channel VIPs.")
    doc.add_paragraph("Features include Telegram Stars backoff, local studio UI, and multi-file processing.")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    upload = UploadFile(filename="release_notes.docx", file=buf)
    photos = await svc.save_and_process_file_upload(upload)

    assert len(photos) >= 1
    p = photos[0]
    assert Path(p.path).exists()
    assert p.source_type == "docx"
    assert p.width == 1200
    assert p.height == 1600


@pytest.mark.asyncio
async def test_process_pdf_upload(tmp_path):
    """Verify PDF document pages are converted to high-resolution JPEG photos."""
    svc = MediaService(media_dir=tmp_path / "media")

    # Generate a 2-page PDF in memory using Pillow
    buf = io.BytesIO()
    img1 = Image.new("RGB", (600, 800), color=(30, 41, 59))
    img2 = Image.new("RGB", (600, 800), color=(51, 65, 85))
    img1.save(buf, format="PDF", save_all=True, append_images=[img2])
    buf.seek(0)

    upload = UploadFile(filename="sample_guide.pdf", file=buf)
    photos = await svc.save_and_process_file_upload(upload)

    assert len(photos) == 2
    for idx, p in enumerate(photos):
        assert Path(p.path).exists()
        assert p.source_type == "pdf"
        assert f"Page {idx + 1}" in p.original_filename
        assert p.size_bytes > 0


@pytest.mark.asyncio
async def test_legacy_save_and_process_upload_compatibility(tmp_path):
    """Verify save_and_process_upload continues to return a single PostPhoto."""
    svc = MediaService(media_dir=tmp_path / "media")

    buf = io.BytesIO()
    img = Image.new("RGB", (500, 500), color="purple")
    img.save(buf, format="JPEG")
    buf.seek(0)

    upload = UploadFile(filename="avatar.jpg", file=buf)
    photo = await svc.save_and_process_upload(upload)

    assert photo is not None
    assert Path(photo.path).exists()
    assert photo.width == 500
    assert photo.source_type == "photo"
