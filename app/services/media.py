import os
import uuid
from pathlib import Path
from typing import List, Optional, Tuple
from fastapi import UploadFile
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from loguru import logger
from app.config import MEDIA_DIR, STATIC_DIR
from app.models.post import PostPhoto

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt"}
ALLOWED_EXTENSIONS = IMAGE_EXTENSIONS | DOCUMENT_EXTENSIONS
MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB per Telegram limits
MAX_DIMENSION = 2560  # Max width/height to avoid Telegram compression artifacts


class MediaService:
    """Service for validating, processing, compressing and storing media files and documents."""

    def __init__(self, media_dir: Path = MEDIA_DIR):
        self.media_dir = media_dir
        self.media_dir.mkdir(parents=True, exist_ok=True)

    def validate_file(self, filename: str, file_size: int) -> Tuple[bool, Optional[str]]:
        """Validate file extension and raw file size."""
        ext = Path(filename).suffix.lower()
        if ext not in ALLOWED_EXTENSIONS:
            return False, f"Unsupported file extension '{ext}'. Allowed: JPG, PNG, WEBP, PDF, DOCX, DOC, TXT."
        if file_size > MAX_FILE_SIZE_BYTES:
            return False, f"File size ({file_size / (1024*1024):.1f}MB) exceeds maximum allowed 10MB."
        return True, None

    def _process_and_save_image(
        self,
        img: Image.Image,
        filename: str,
        unique_suffix: str = "",
        source_type: str = "photo",
        apply_watermark_hook: Optional[callable] = None,
    ) -> PostPhoto:
        """Helper to normalize mode, resize, watermark, and save a PIL Image as JPEG."""
        unique_id = uuid.uuid4().hex[:12]
        suffix_str = f"_{unique_suffix}" if unique_suffix else ""
        safe_name = f"{unique_id}{suffix_str}.jpg"
        dest_path = self.media_dir / safe_name

        # Convert to RGB if RGBA, LA, or P to save as clean JPEG
        if img.mode in ("RGBA", "LA", "P"):
            bg_color = (15, 23, 42) if source_type != "photo" else (0, 0, 0)
            background = Image.new("RGB", img.size, bg_color)
            if img.mode == "P":
                img = img.convert("RGBA")
            background.paste(img, mask=img.split()[-1] if "A" in img.mode else None)
            img = background
        elif img.mode != "RGB":
            img = img.convert("RGB")

        # Resize if dimensions exceed Telegram recommendation
        orig_w, orig_h = img.size
        if orig_w > MAX_DIMENSION or orig_h > MAX_DIMENSION:
            img.thumbnail((MAX_DIMENSION, MAX_DIMENSION), Image.Resampling.LANCZOS)

        # Optional watermark hook from plugin
        if apply_watermark_hook:
            try:
                img = apply_watermark_hook(img)
            except Exception as w_err:
                logger.warning(f"Watermark hook failed, continuing without watermark: {w_err}")

        # Save optimized JPEG
        img.save(dest_path, "JPEG", quality=88, optimize=True)
        final_size = os.path.getsize(dest_path)
        w, h = img.size

        logger.info(f"Saved processed {source_type} media to {dest_path} ({w}x{h}, {final_size} bytes)")

        return PostPhoto(
            path=str(dest_path),
            original_filename=filename,
            width=w,
            height=h,
            size_bytes=final_size,
            source_type=source_type,
        )

    def _render_document_text_to_cards(
        self,
        title: str,
        text: str,
        doc_type: str = "DOC"
    ) -> List[Image.Image]:
        """
        Renders long document text into styled, publication-ready 1200x1600 card slides.
        Splits text across multiple page cards (max 10).
        """
        width = 1200
        height = 1600
        margin = 80
        content_w = width - (margin * 2)

        badge_font = self._load_font(22)
        title_font = self._load_font(30)
        body_font = self._load_font(26)
        footer_font = self._load_font(20)

        # Wrap text lines
        wrapped_lines = self._wrap_text(text, body_font, content_w)
        lines_per_page = 26
        line_height = 42

        # Break lines into pages
        pages_lines = []
        for i in range(0, max(1, len(wrapped_lines)), lines_per_page):
            pages_lines.append(wrapped_lines[i : i + lines_per_page])

        total_pages = min(len(pages_lines), 10)
        card_images = []

        for p_idx in range(total_pages):
            page_num = p_idx + 1
            cur_lines = pages_lines[p_idx]

            # Background gradient
            img = Image.new("RGB", (width, height))
            top_color = (15, 23, 42)    # Slate 900
            bottom_color = (30, 27, 75) # Indigo 950
            for y in range(height):
                ratio = y / max(1, height - 1)
                r = int(top_color[0] + (bottom_color[0] - top_color[0]) * ratio)
                g = int(top_color[1] + (bottom_color[1] - top_color[1]) * ratio)
                b = int(top_color[2] + (bottom_color[2] - top_color[2]) * ratio)
                for x in range(width):
                    img.putpixel((x, y), (r, g, b))

            draw = ImageDraw.Draw(img)

            # Elegant border
            draw.rounded_rectangle(
                [25, 25, width - 25, height - 25],
                radius=24,
                outline=(255, 255, 255, 35),
                width=2
            )

            # Header Badge
            badge_text = f"★ {doc_type} DOCUMENT"
            draw.text((margin, 60), badge_text, fill=(245, 158, 11), font=badge_font)

            # Page count badge
            page_text = f"PAGE {page_num} OF {total_pages}"
            try:
                page_bbox = badge_font.getbbox(page_text)
                page_w = page_bbox[2] - page_bbox[0]
            except Exception:
                page_w = len(page_text) * 12
            draw.text((width - margin - page_w, 60), page_text, fill=(148, 163, 184), font=badge_font)

            # Document Title
            display_title = (title[:48] + "...") if len(title) > 50 else title
            draw.text((margin, 100), display_title, fill=(255, 255, 255), font=title_font)

            # Header separator line
            draw.line([margin, 155, width - margin, 155], fill=(255, 255, 255, 40), width=1)

            # Body text lines
            y_offset = 185
            for line in cur_lines:
                draw.text((margin, y_offset), line, fill=(226, 232, 240), font=body_font)
                y_offset += line_height

            # Footer separator line
            draw.line([margin, height - 90, width - margin, height - 90], fill=(255, 255, 255, 30), width=1)

            # Footer
            draw.text((margin, height - 70), "★ Telegram Stars Paid Content • StarPost Studio", fill=(100, 116, 139), font=footer_font)
            card_images.append(img)

        return card_images

    def _convert_pdf_to_images(self, content: bytes, original_filename: str) -> List[Tuple[Image.Image, str]]:
        """Render PDF pages into PIL Images using pypdfium2."""
        import pypdfium2
        pdf = pypdfium2.PdfDocument(content)
        page_count = min(len(pdf), 10)
        results = []
        for idx in range(page_count):
            page = pdf[idx]
            pil_img = page.render(scale=2.0).to_pil()
            page_title = f"{original_filename} (Page {idx + 1})"
            results.append((pil_img, page_title))
        return results

    def _convert_docx_to_images(self, content: bytes, original_filename: str) -> List[Tuple[Image.Image, str]]:
        """Extract embedded images and formatted text slides from DOCX."""
        import io
        import docx
        doc = docx.Document(io.BytesIO(content))
        results = []

        # 1. Extract embedded images
        img_idx = 1
        for rel in doc.part.rels.values():
            if "image" in getattr(rel, "target_ref", ""):
                try:
                    img_data = rel.target_part.blob
                    pil_img = Image.open(io.BytesIO(img_data))
                    results.append((pil_img, f"{original_filename} (Image {img_idx})"))
                    img_idx += 1
                except Exception:
                    pass

        # 2. Extract text paragraphs
        text_lines = [p.text for p in doc.paragraphs if p.text.strip()]
        full_text = "\n\n".join(text_lines).strip()

        if full_text:
            text_cards = self._render_document_text_to_cards(
                title=Path(original_filename).stem,
                text=full_text,
                doc_type="DOCX"
            )
            for p_idx, card_img in enumerate(text_cards, start=1):
                results.append((card_img, f"{original_filename} (Page {p_idx})"))

        if not results:
            empty_cards = self._render_document_text_to_cards(
                title=Path(original_filename).stem,
                text="(Empty Document)",
                doc_type="DOCX"
            )
            results.append((empty_cards[0], original_filename))

        return results[:10]

    def _convert_doc_to_images(self, content: bytes, original_filename: str) -> List[Tuple[Image.Image, str]]:
        """Extract text from legacy DOC or DOCX and render into slides."""
        try:
            return self._convert_docx_to_images(content, original_filename)
        except Exception:
            import re
            text_chunks = re.findall(rb'[\x20-\x7E\t\r\n]{4,}', content)
            text = "\n".join([chunk.decode('latin-1', errors='ignore') for chunk in text_chunks])
            if not text.strip():
                text = f"(Document content extracted from {original_filename})"
            text_cards = self._render_document_text_to_cards(
                title=Path(original_filename).stem,
                text=text,
                doc_type="DOC"
            )
            return [(img, f"{original_filename} (Page {i+1})") for i, img in enumerate(text_cards[:10])]

    def _convert_txt_to_images(self, content: bytes, original_filename: str) -> List[Tuple[Image.Image, str]]:
        """Decode and render text files into formatted slides."""
        text = ""
        for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                text = content.decode(enc)
                break
            except Exception:
                continue
        if not text.strip():
            text = "(Empty Text File)"
        text_cards = self._render_document_text_to_cards(
            title=Path(original_filename).stem,
            text=text,
            doc_type="TXT"
        )
        return [(img, f"{original_filename} (Page {i+1})") for i, img in enumerate(text_cards[:10])]

    async def save_and_process_file_upload(
        self,
        upload_file: UploadFile,
        apply_watermark_hook: Optional[callable] = None
    ) -> List[PostPhoto]:
        """
        Process any supported file: JPG/PNG/WEBP, PDF, DOCX/DOC, or TXT.
        Images are processed directly. Documents are rendered into high-resolution
        publication-ready Telegram photo slides (up to 10 pages).
        Returns a list of PostPhoto instances.
        """
        filename = upload_file.filename or "file.jpg"
        content = await upload_file.read()
        file_size = len(content)

        valid, error_msg = self.validate_file(filename, file_size)
        if not valid:
            raise ValueError(error_msg)

        ext = Path(filename).suffix.lower()

        try:
            if ext in IMAGE_EXTENSIONS:
                from io import BytesIO
                img = Image.open(BytesIO(content))
                photo = self._process_and_save_image(
                    img=img,
                    filename=filename,
                    source_type="photo",
                    apply_watermark_hook=apply_watermark_hook,
                )
                return [photo]

            elif ext == ".pdf":
                raw_images = self._convert_pdf_to_images(content, filename)
                photos = []
                for idx, (img, page_title) in enumerate(raw_images):
                    photo = self._process_and_save_image(
                        img=img,
                        filename=page_title,
                        unique_suffix=f"p{idx+1}",
                        source_type="pdf",
                        apply_watermark_hook=apply_watermark_hook,
                    )
                    photos.append(photo)
                return photos

            elif ext == ".docx":
                raw_images = self._convert_docx_to_images(content, filename)
                photos = []
                for idx, (img, page_title) in enumerate(raw_images):
                    photo = self._process_and_save_image(
                        img=img,
                        filename=page_title,
                        unique_suffix=f"doc_p{idx+1}",
                        source_type="docx",
                        apply_watermark_hook=apply_watermark_hook,
                    )
                    photos.append(photo)
                return photos

            elif ext == ".doc":
                raw_images = self._convert_doc_to_images(content, filename)
                photos = []
                for idx, (img, page_title) in enumerate(raw_images):
                    photo = self._process_and_save_image(
                        img=img,
                        filename=page_title,
                        unique_suffix=f"doc_p{idx+1}",
                        source_type="doc",
                        apply_watermark_hook=apply_watermark_hook,
                    )
                    photos.append(photo)
                return photos

            elif ext == ".txt":
                raw_images = self._convert_txt_to_images(content, filename)
                photos = []
                for idx, (img, page_title) in enumerate(raw_images):
                    photo = self._process_and_save_image(
                        img=img,
                        filename=page_title,
                        unique_suffix=f"txt_p{idx+1}",
                        source_type="txt",
                        apply_watermark_hook=apply_watermark_hook,
                    )
                    photos.append(photo)
                return photos

            else:
                raise ValueError(f"Unsupported file format: {ext}")

        except Exception as exc:
            logger.error(f"Media processing failed for {filename}: {exc}")
            raise ValueError(f"Failed to process {filename}: {exc}")

    async def save_and_process_upload(
        self,
        upload_file: UploadFile,
        apply_watermark_hook: Optional[callable] = None
    ) -> PostPhoto:
        """Process an uploaded file, returning the primary/first PostPhoto."""
        photos = await self.save_and_process_file_upload(upload_file, apply_watermark_hook)
        if not photos:
            raise ValueError(f"No photos generated from {upload_file.filename}")
        return photos[0]

    def delete_photo(self, photo_path: str) -> bool:
        """Safely delete photo from disk."""
        try:
            path = Path(photo_path).resolve()
            if path.exists():
                path.unlink()
                logger.info(f"Deleted media file {path}")
                return True
            return False
        except Exception as exc:
            logger.warning(f"Failed to delete photo {photo_path}: {exc}")
            return False

    def _load_font(self, size: int) -> ImageFont.ImageFont:
        """Load bundled DejaVuSans or NotoSans TTF font, falling back gracefully to default."""
        font_candidates = [
            STATIC_DIR / "fonts" / "DejaVuSans.ttf",
            STATIC_DIR / "fonts" / "NotoSans-Regular.ttf",
        ]
        for fp in font_candidates:
            if fp.exists():
                try:
                    return ImageFont.truetype(str(fp), size)
                except Exception:
                    continue
        return ImageFont.load_default()

    def _wrap_text(self, text: str, font: ImageFont.ImageFont, max_width: int) -> List[str]:
        """Wrap text to fit within max_width pixels."""
        lines = []
        for paragraph in text.split("\n"):
            p_strip = paragraph.strip()
            if not p_strip:
                lines.append("")
                continue
            words = p_strip.split(" ")
            cur = []
            for word in words:
                candidate = " ".join(cur + [word])
                try:
                    bbox = font.getbbox(candidate)
                    w = bbox[2] - bbox[0]
                except Exception:
                    w = len(candidate) * getattr(font, "size", 16) * 0.55
                if w <= max_width:
                    cur.append(word)
                else:
                    if cur:
                        lines.append(" ".join(cur))
                        cur = [word]
                    else:
                        lines.append(word)
                        cur = []
            if cur:
                lines.append(" ".join(cur))
        return lines

    def render_text_card(
        self,
        text: str,
        reference_size: Tuple[int, int],
        style: str = "blurred photo background",
        background_photo_path: Optional[str] = None,
    ) -> str:
        """
        Draws hidden text, word-wrapped and auto-sized, on a card the same size as reference_size.
        Styles: 'blurred photo background', 'gradient', 'solid'.
        Returns: file path to rendered card JPEG.
        """
        ref_w, ref_h = reference_size
        width = max(600, ref_w)
        height = max(600, ref_h)

        # 1. Base card background
        style_normalized = (style or "blurred photo background").lower().strip()
        card_img = None
        if ("blur" in style_normalized or "photo" in style_normalized) and background_photo_path and Path(background_photo_path).exists():
            try:
                bg = Image.open(background_photo_path).convert("RGB")
                bg = bg.resize((width, height), Image.Resampling.LANCZOS)
                bg = bg.filter(ImageFilter.GaussianBlur(radius=35))
                darkener = Image.new("RGBA", (width, height), (15, 23, 42, 195))
                card_img = Image.alpha_composite(bg.convert("RGBA"), darkener).convert("RGB")
            except Exception as e:
                logger.warning(f"Could not blur background photo for text card: {e}")
                card_img = None

        if card_img is None:
            if "solid" in style_normalized:
                card_img = Image.new("RGB", (width, height), (15, 23, 42))
            else:
                # Gradient (from slate-900 to indigo-950)
                card_img = Image.new("RGB", (width, height))
                top_color = (15, 23, 42)
                bottom_color = (30, 27, 75)
                for y in range(height):
                    ratio = y / max(1, height - 1)
                    r = int(top_color[0] + (bottom_color[0] - top_color[0]) * ratio)
                    g = int(top_color[1] + (bottom_color[1] - top_color[1]) * ratio)
                    b = int(top_color[2] + (bottom_color[2] - top_color[2]) * ratio)
                    for x in range(width):
                        card_img.putpixel((x, y), (r, g, b))

        draw = ImageDraw.Draw(card_img)

        # Decorative inner border
        border_inset = int(min(width, height) * 0.04)
        draw.rounded_rectangle(
            [border_inset, border_inset, width - border_inset, height - border_inset],
            radius=20,
            outline=(255, 255, 255, 40),
            width=2,
        )

        margin = int(min(width, height) * 0.08)
        content_w = width - 2 * margin
        header_y = margin + 10
        footer_space = int(min(width, height) * 0.12)
        content_h = height - header_y - footer_space - 20

        # Header Badge
        header_font = self._load_font(max(14, int(min(width, height) * 0.028)))
        badge_text = "★ UNLOCKED BUYER CONTENT"
        draw.text((margin, header_y), badge_text, fill=(245, 158, 11), font=header_font)
        draw.line([margin, header_y + 35, width - margin, header_y + 35], fill=(255, 255, 255, 30), width=1)

        # 2. Auto-size and word-wrap hidden text
        text_content_y = header_y + 55
        target_size = min(max(20, int(min(width, height) * 0.038)), 46)
        min_size = 14
        chosen_font = self._load_font(target_size)
        wrapped_lines = self._wrap_text(text, chosen_font, content_w)

        while target_size > min_size:
            line_spacing = int(target_size * 1.45)
            total_h = len(wrapped_lines) * line_spacing
            if total_h <= content_h:
                break
            target_size -= 2
            chosen_font = self._load_font(target_size)
            wrapped_lines = self._wrap_text(text, chosen_font, content_w)

        line_spacing = int(target_size * 1.45)

        # Render lines
        curr_y = text_content_y
        for line in wrapped_lines:
            if curr_y + line_spacing > height - footer_space:
                draw.text((margin, curr_y), "...", fill=(248, 250, 252), font=chosen_font)
                break
            draw.text((margin, curr_y), line, fill=(248, 250, 252), font=chosen_font)
            curr_y += line_spacing

        # Footer Badge
        footer_font = self._load_font(max(12, int(min(width, height) * 0.022)))
        footer_text = "Delivered via Telegram Stars Paid Media"
        draw.text((margin, height - margin - 15), footer_text, fill=(148, 163, 184), font=footer_font)

        # Save card
        unique_id = uuid.uuid4().hex[:12]
        dest_path = self.media_dir / f"card_{unique_id}.jpg"
        card_img.save(dest_path, "JPEG", quality=90, optimize=True)
        logger.info(f"Generated hidden text card at {dest_path} ({width}x{height})")
        return str(dest_path)

    def generate_teaser(self, photo_path: str, star_count: int) -> str:
        """
        Create a blurred teaser image of photo 1 with an 'Unlock for N Stars' badge overlay.
        Returns: file path to teaser JPEG.
        """
        path = Path(photo_path)
        if not path.exists():
            raise FileNotFoundError(f"Media file not found: {photo_path}")

        img = Image.open(path).convert("RGB")
        w, h = img.size

        # Heavy Gaussian blur
        teaser = img.filter(ImageFilter.GaussianBlur(radius=30))

        # Darken overlay
        overlay = Image.new("RGBA", (w, h), (0, 0, 0, 95))
        teaser = Image.alpha_composite(teaser.convert("RGBA"), overlay).convert("RGB")

        # Center Unlock badge
        badge_w = int(w * 0.65)
        badge_h = int(h * 0.22)
        badge_x0 = (w - badge_w) // 2
        badge_y0 = (h - badge_h) // 2
        badge_x1 = badge_x0 + badge_w
        badge_y1 = badge_y0 + badge_h

        # Translucent pill background
        pill = Image.new("RGBA", (w, h), (0, 0, 0, 0))
        pill_draw = ImageDraw.Draw(pill)
        pill_draw.rounded_rectangle(
            [badge_x0, badge_y0, badge_x1, badge_y1],
            radius=int(badge_h * 0.35),
            fill=(15, 23, 42, 210),
            outline=(245, 158, 11, 230),
            width=3,
        )
        teaser = Image.alpha_composite(teaser.convert("RGBA"), pill).convert("RGB")

        # Text inside badge
        draw = ImageDraw.Draw(teaser)
        title_font_size = max(18, int(badge_h * 0.30))
        title_font = self._load_font(title_font_size)
        sub_font_size = max(12, int(badge_h * 0.18))
        sub_font = self._load_font(sub_font_size)

        title_text = f"★ Unlock for {star_count} Stars"
        sub_text = "Locked content • Tap link below to purchase"

        try:
            tb = title_font.getbbox(title_text)
            tw = tb[2] - tb[0]
            th = tb[3] - tb[1]
        except Exception:
            tw, th = len(title_text) * 12, title_font_size

        try:
            sb = sub_font.getbbox(sub_text)
            sw = sb[2] - sb[0]
        except Exception:
            sw = len(sub_text) * 8

        tx = (w - tw) // 2
        ty = badge_y0 + int(badge_h * 0.22)
        draw.text((tx, ty), title_text, fill=(245, 158, 11), font=title_font)

        sx = (w - sw) // 2
        sy = ty + th + int(badge_h * 0.15)
        draw.text((sx, sy), sub_text, fill=(248, 250, 252), font=sub_font)

        # Save teaser
        unique_id = uuid.uuid4().hex[:12]
        dest_path = self.media_dir / f"teaser_{unique_id}.jpg"
        teaser.save(dest_path, "JPEG", quality=88, optimize=True)
        logger.info(f"Generated teaser image at {dest_path} ({w}x{h})")
        return str(dest_path)


media_service = MediaService()
